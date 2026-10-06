import pandas as pd
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation


# ---------------------------------------------------------
# PATHS
# ---------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent
STEP2_DIR = BASE_DIR.parent / "step2"

OUTPUT_FILE = BASE_DIR / "labeling_dataset.xlsx"


# ---------------------------------------------------------
# LOAD THE THREE FRESH DATASETS
# ---------------------------------------------------------

datasets = [
    ("LlamaIndex", STEP2_DIR / "pr_llamaindex.xlsx"),
    ("Haystack", STEP2_DIR / "pr_haystack.xlsx"),
    ("LangChain", STEP2_DIR / "pr_langchain.xlsx"),
]

frames = []

for framework, path in datasets:

    df = pd.read_excel(path)

    df.insert(0, "Framework", framework)

    frames.append(df)


all_prs = pd.concat(frames, ignore_index=True)

print("Total PRs loaded:", len(all_prs))

print(
    all_prs["Framework"]
    .value_counts()
    .to_string()
)


# ---------------------------------------------------------
# LABELING COLUMNS
# ---------------------------------------------------------

label_columns = [

    # Annotator 1
    "A1_Bug_Status",
    "A1_Root_Cause",
    "A1_Symptom",
    "A1_Component",
    "A1_Notes",

    # Annotator 2
    "A2_Bug_Status",
    "A2_Root_Cause",
    "A2_Symptom",
    "A2_Component",
    "A2_Notes",

    # Consensus / final classification
    "Final_Bug_Status",
    "Final_Root_Cause",
    "Final_Symptom",
    "Final_Component",
    "Final_Notes",
]

for column in label_columns:
    all_prs[column] = ""


# ---------------------------------------------------------
# RANDOMIZE DATASET
#
# Paper states random 30% pilot sample.
# The paper does not report the random seed.
# We use seed 42 so our reproduction is repeatable.
# ---------------------------------------------------------

all_prs = all_prs.sample(
    frac=1,
    random_state=42
).reset_index(drop=True)


# ---------------------------------------------------------
# 30% PILOT SAMPLE
# ---------------------------------------------------------

pilot_size = round(len(all_prs) * 0.30)

pilot = all_prs.iloc[:pilot_size].copy()

remaining = all_prs.iloc[pilot_size:].copy()
remaining = remaining.reset_index(drop=True)


print("\nPilot sample:", len(pilot))
print("Remaining:", len(remaining))


# ---------------------------------------------------------
# SPLIT REMAINING 70% INTO FOUR ROUNDS
# ---------------------------------------------------------

base_size = len(remaining) // 4
remainder = len(remaining) % 4

round_sizes = [
    base_size + (1 if i < remainder else 0)
    for i in range(4)
]

rounds = []

start = 0

for size in round_sizes:

    end = start + size

    rounds.append(
        remaining.iloc[start:end].copy()
    )

    start = end


for i, round_df in enumerate(rounds, start=1):

    print(
        f"Round {i}: {len(round_df)} PRs"
    )


# ---------------------------------------------------------
# TAXONOMY
# ---------------------------------------------------------

taxonomy = [

    # Root Causes
    [
        "Root Cause",
        "R1",
        "API Misuse",
        "Incorrect, missing, wrong, or redundant API usage."
    ],
    [
        "Root Cause",
        "R2",
        "Incompatibility",
        "Functional or version incompatibility between modules, libraries, or dependencies."
    ],
    [
        "Root Cause",
        "R3",
        "Assignment Issue",
        "Incorrect assignment, initialization, value, type, or missing initialization."
    ],
    [
        "Root Cause",
        "R4",
        "Parameter/Argument Issue",
        "Missing, incorrect, incomplete, redundant, or conflicting parameters or arguments."
    ],
    [
        "Root Cause",
        "R5",
        "Code Logic Issue",
        "Incorrect program logic, workflow, conditions, ordering, or algorithm design."
    ],
    [
        "Root Cause",
        "R6",
        "Import Error",
        "Missing, incorrect, or redundant internal or external imports."
    ],
    [
        "Root Cause",
        "R7",
        "Typo",
        "Typographical mistakes in syntax, names, variables, or function calls."
    ],
    [
        "Root Cause",
        "R8",
        "Incorrect Exception Handling",
        "Exceptions are not handled correctly, producing misleading errors or silent failures."
    ],
    [
        "Root Cause",
        "R9",
        "Incorrect Numerical Computation",
        "Incorrect calculations, rounding, precision, or mathematical operations."
    ],

    # Symptoms
    [
        "Symptom",
        "S1",
        "Crash",
        "Runtime execution terminates or throws an exception."
    ],
    [
        "Symptom",
        "S2",
        "Incorrect Functionality",
        "Program runs but does not perform its intended function correctly."
    ],
    [
        "Symptom",
        "S3",
        "Unexpected Output",
        "Program produces output that deviates unexpectedly from expected results."
    ],
    [
        "Symptom",
        "S4",
        "Hang",
        "Program stops progressing but does not terminate."
    ],
    [
        "Symptom",
        "S5",
        "External Connection Failure",
        "Failure involving external services, libraries, APIs, resources, or dependencies."
    ],
    [
        "Symptom",
        "S6",
        "Unidentified",
        "Available PR evidence is insufficient to identify a specific symptom."
    ],

    # Components
    [
        "Component",
        "DP",
        "Data Preprocessing",
        "Loaders, embeddings, retrievers, text processing, indexing, and vector/data preparation."
    ],
    [
        "Component",
        "CS",
        "Core Schema",
        "Prompt templates, chat/model interfaces, and output parsers."
    ],
    [
        "Component",
        "AC",
        "Agent Construction",
        "Agent creation, tool selection, orchestration, and dynamic agent workflows."
    ],
    [
        "Component",
        "FM",
        "Featured Module",
        "Framework-specific specialized modules and advanced functionality."
    ],
]

taxonomy_df = pd.DataFrame(
    taxonomy,
    columns=[
        "Dimension",
        "Code",
        "Category",
        "Description"
    ]
)


# ---------------------------------------------------------
# COLLECTION SUMMARY
# ---------------------------------------------------------

summary = pd.DataFrame(
    [
        ["Total fresh PRs", len(all_prs)],
        [
            "LlamaIndex PRs",
            len(
                all_prs[
                    all_prs["Framework"]
                    == "LlamaIndex"
                ]
            )
        ],
        [
            "Haystack PRs",
            len(
                all_prs[
                    all_prs["Framework"]
                    == "Haystack"
                ]
            )
        ],
        [
            "LangChain PRs",
            len(
                all_prs[
                    all_prs["Framework"]
                    == "LangChain"
                ]
            )
        ],
        ["Pilot 30%", len(pilot)],
        ["Remaining 70%", len(remaining)],
        ["Round 1", len(rounds[0])],
        ["Round 2", len(rounds[1])],
        ["Round 3", len(rounds[2])],
        ["Round 4", len(rounds[3])],
        ["Random seed", 42],
    ],
    columns=[
        "Metric",
        "Value"
    ]
)


# ---------------------------------------------------------
# WRITE WORKBOOK
# ---------------------------------------------------------

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    summary.to_excel(
        writer,
        sheet_name="Collection_Summary",
        index=False
    )

    pilot.to_excel(
        writer,
        sheet_name="Pilot_30pct",
        index=False
    )

    for i, round_df in enumerate(rounds, start=1):

        round_df.to_excel(
            writer,
            sheet_name=f"Round{i}",
            index=False
        )

    taxonomy_df.to_excel(
        writer,
        sheet_name="Taxonomy",
        index=False
    )


# ---------------------------------------------------------
# FORMAT WORKBOOK
# ---------------------------------------------------------

wb = load_workbook(OUTPUT_FILE)

header_fill = PatternFill(
    "solid",
    fgColor="1F4E78"
)

header_font = Font(
    color="FFFFFF",
    bold=True
)

thin_border = Border(
    bottom=Side(
        style="thin",
        color="D9E1F2"
    )
)


label_sheets = [
    "Pilot_30pct",
    "Round1",
    "Round2",
    "Round3",
    "Round4"
]


# ---------------------------------------------------------
# DROPDOWN VALUES
# ---------------------------------------------------------

bug_status_values = [
    "Include",
    "Exclude",
    "Unsure"
]

root_cause_values = [
    "R1 API Misuse",
    "R2 Incompatibility",
    "R3 Assignment Issue",
    "R4 Parameter/Argument Issue",
    "R5 Code Logic Issue",
    "R6 Import Error",
    "R7 Typo",
    "R8 Incorrect Exception Handling",
    "R9 Incorrect Numerical Computation"
]

symptom_values = [
    "S1 Crash",
    "S2 Incorrect Functionality",
    "S3 Unexpected Output",
    "S4 Hang",
    "S5 External Connection Failure",
    "S6 Unidentified"
]

component_values = [
    "DP Data Preprocessing",
    "CS Core Schema",
    "AC Agent Construction",
    "FM Featured Module"
]


# ---------------------------------------------------------
# APPLY FORMATTING + VALIDATION
# ---------------------------------------------------------

for sheet_name in label_sheets:

    ws = wb[sheet_name]

    ws.freeze_panes = "E2"

    ws.auto_filter.ref = ws.dimensions

    ws.sheet_view.showGridLines = False

    # Header formatting
    for cell in ws[1]:

        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    # Widths
    widths = {
        "A": 14,
        "B": 10,
        "C": 60,
        "D": 55,

        "E": 18,
        "F": 30,
        "G": 28,
        "H": 25,
        "I": 40,

        "J": 18,
        "K": 30,
        "L": 28,
        "M": 25,
        "N": 40,

        "O": 20,
        "P": 30,
        "Q": 28,
        "R": 25,
        "S": 45,
    }

    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    # Wrap text
    for row in ws.iter_rows(
        min_row=2,
        max_row=ws.max_row
    ):

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

            cell.border = thin_border

    # Make GitHub URLs clickable
    for row in range(2, ws.max_row + 1):

        url_cell = ws[f"D{row}"]

        if url_cell.value:

            url_cell.hyperlink = url_cell.value
            url_cell.style = "Hyperlink"

    # Bug status dropdown
    dv_bug = DataValidation(
        type="list",
        formula1='"'
        + ",".join(bug_status_values)
        + '"',
        allow_blank=True
    )

    # Root cause dropdown
    dv_root = DataValidation(
        type="list",
        formula1='"'
        + ",".join(root_cause_values)
        + '"',
        allow_blank=True
    )

    # Symptom dropdown
    dv_symptom = DataValidation(
        type="list",
        formula1='"'
        + ",".join(symptom_values)
        + '"',
        allow_blank=True
    )

    # Component dropdown
    dv_component = DataValidation(
        type="list",
        formula1='"'
        + ",".join(component_values)
        + '"',
        allow_blank=True
    )

    ws.add_data_validation(dv_bug)
    ws.add_data_validation(dv_root)
    ws.add_data_validation(dv_symptom)
    ws.add_data_validation(dv_component)

    max_row = ws.max_row

    # Annotator 1
    dv_bug.add(f"E2:E{max_row}")
    dv_root.add(f"F2:F{max_row}")
    dv_symptom.add(f"G2:G{max_row}")
    dv_component.add(f"H2:H{max_row}")

    # Annotator 2
    dv_bug.add(f"J2:J{max_row}")
    dv_root.add(f"K2:K{max_row}")
    dv_symptom.add(f"L2:L{max_row}")
    dv_component.add(f"M2:M{max_row}")

    # Final consensus
    dv_bug.add(f"O2:O{max_row}")
    dv_root.add(f"P2:P{max_row}")
    dv_symptom.add(f"Q2:Q{max_row}")
    dv_component.add(f"R2:R{max_row}")


# ---------------------------------------------------------
# TAXONOMY FORMATTING
# ---------------------------------------------------------

taxonomy_ws = wb["Taxonomy"]

taxonomy_ws.freeze_panes = "A2"

taxonomy_ws.sheet_view.showGridLines = False

for cell in taxonomy_ws[1]:

    cell.fill = header_fill
    cell.font = header_font

taxonomy_ws.column_dimensions["A"].width = 18
taxonomy_ws.column_dimensions["B"].width = 10
taxonomy_ws.column_dimensions["C"].width = 35
taxonomy_ws.column_dimensions["D"].width = 90

for row in taxonomy_ws.iter_rows():

    for cell in row:

        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )


# ---------------------------------------------------------
# SUMMARY FORMATTING
# ---------------------------------------------------------

summary_ws = wb["Collection_Summary"]

for cell in summary_ws[1]:

    cell.fill = header_fill
    cell.font = header_font

summary_ws.column_dimensions["A"].width = 28
summary_ws.column_dimensions["B"].width = 15


# ---------------------------------------------------------
# SAVE
# ---------------------------------------------------------

wb.save(OUTPUT_FILE)

print("\nLabeling workbook created successfully.")
print("Saved to:", OUTPUT_FILE)